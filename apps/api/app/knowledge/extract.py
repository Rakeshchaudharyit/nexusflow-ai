"""Upload validation and text extraction for the four supported formats.

TXT/Markdown are decoded as UTF-8; PDF text comes from pypdf (no OCR: a scanned PDF fails as
`no_text`); DOCX paragraphs and tables come from python-docx. Uploaded content is only parsed,
never executed, and filenames are display names only.
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

from app.agents.errors import AgentRuntimeError

MIME_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
# Declared types we accept per extension (browsers disagree about Markdown).
DECLARED = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/x-markdown", "text/plain"},
    ".pdf": {"application/pdf"},
    ".docx": {MIME_TYPES[".docx"]},
}
GENERIC = {"", "application/octet-stream"}
SAFE_NAME = re.compile(r"^[\w][\w .,()&'+\-]*$")
MAX_PDF_PAGES = 1000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MIN_TEXT_CHARS = 20


class KnowledgeFileError(AgentRuntimeError):
    """Permanent: the file itself is unacceptable. Safe code and message only."""

    code = "invalid_file"
    message = "The file could not be read."


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None


@dataclass(frozen=True)
class Extraction:
    sections: list[Section]
    parser: str
    pages: int | None = None
    chars: int = field(default=0)


def safe_filename(name: str) -> str:
    """Validate a display name. Anything that looks like a path is rejected, not cleaned."""
    name = unicodedata.normalize("NFC", (name or "").strip())
    if (
        not name
        or len(name) > 200
        or "/" in name
        or "\\" in name
        or ".." in name
        or any(unicodedata.category(char).startswith("C") for char in name)
        or not SAFE_NAME.match(name)
    ):
        raise KnowledgeFileError(
            "The filename is not allowed. Use letters, numbers, spaces and . _ - ( ) only.",
            code="unsafe_filename",
        )
    return name


def file_type(name: str, declared: str | None) -> tuple[str, str]:
    """Return (extension, canonical MIME type) or raise unsupported_file."""
    extension = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if extension not in MIME_TYPES:
        raise KnowledgeFileError(
            "Unsupported file type. Upload .txt, .md, .pdf or .docx.", code="unsupported_file"
        )
    declared = (declared or "").split(";")[0].strip().lower()
    if declared not in GENERIC and declared not in DECLARED[extension]:
        raise KnowledgeFileError(
            "The file's content type does not match its extension.", code="unsupported_file"
        )
    return extension, MIME_TYPES[extension]


def check_content(extension: str, data: bytes) -> None:
    """Cheap signature checks before anything is stored or queued."""
    if not data:
        raise KnowledgeFileError("The file is empty.", code="invalid_file")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise KnowledgeFileError("The file is not a PDF.", code="invalid_file")
    if extension == ".docx" and not data.startswith(b"PK\x03\x04"):
        raise KnowledgeFileError("The file is not a DOCX document.", code="invalid_file")
    if extension in (".txt", ".md"):
        decode_text(data)


def decode_text(data: bytes) -> str:
    if b"\x00" in data:
        raise KnowledgeFileError("Text files must not contain binary data.", code="invalid_file")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise KnowledgeFileError("Text files must be UTF-8.", code="invalid_file") from error


def _pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise KnowledgeFileError("Encrypted PDFs are not supported.", code="unsupported_file")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise KnowledgeFileError("The PDF has too many pages.", code="invalid_file")
        sections = [
            Section((page.extract_text() or "").strip(), number)
            for number, page in enumerate(reader.pages, start=1)
        ]
    except KnowledgeFileError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as error:
        raise KnowledgeFileError("The PDF could not be read.", code="invalid_file") from error
    return Extraction([s for s in sections if s.text], "pypdf", pages=len(sections))


def _docx(data: bytes) -> Extraction:
    import docx

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise KnowledgeFileError("The DOCX expands too much.", code="invalid_file")
            if "word/document.xml" not in archive.namelist():
                raise KnowledgeFileError("The file is not a DOCX document.", code="invalid_file")
        document = docx.Document(io.BytesIO(data))
    except KnowledgeFileError:
        raise
    except Exception as error:  # python-docx/lxml raise many types for malformed input
        raise KnowledgeFileError("The DOCX could not be read.", code="invalid_file") from error
    lines = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading") or style.lower() == "title":
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lines += ["", "#" * min(level, 6) + " " + text, ""]
        else:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Extraction([Section("\n".join(lines).strip())], "python-docx")


def extract(extension: str, data: bytes) -> Extraction:
    if extension == ".pdf":
        extraction = _pdf(data)
    elif extension == ".docx":
        extraction = _docx(data)
    else:
        text = decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
        extraction = Extraction([Section(text.strip())], "text")
    chars = sum(len(section.text) for section in extraction.sections)
    if chars < MIN_TEXT_CHARS:
        raise KnowledgeFileError(
            "No extractable text was found. Scanned documents need OCR, which is not supported.",
            code="no_text",
        )
    return Extraction(extraction.sections, extraction.parser, extraction.pages, chars)
